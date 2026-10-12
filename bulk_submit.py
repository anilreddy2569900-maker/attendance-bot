"""
Headless bulk submitter for the Apps Script "Attendance" form.

Usage:
    python bulk_submit.py students.xlsx                  # normal use (run.bat does this)
    python bulk_submit.py students.xlsx --otp 123456     # no prompt, single round
    python bulk_submit.py students.xlsx --mode click     # type into the boxes + click instead

Flow: reads the sheet -> asks for the OTP right away while the form loads in the
background -> submits everyone -> shows OK / FAIL / ERROR / CHECK per student ->
lets you retry just the failed ones -> writes results_<time>.csv.
"""
import argparse
import csv
import queue
import re
import sys
import threading
import time
from pathlib import Path

from playwright.sync_api import sync_playwright

URL = "https://script.google.com/macros/s/AKfycbzDBQiJ1ibYdBvFMAtZPU-YfEoWlFAxybctQ8KU_trhm3GX4ZYoamIWlMl7gMz1mS5l/exec"
SERVER_FN = "submitStudentAttendance"   # what the form's button calls
LOAD_TIMEOUT_MS = 45_000
CALL_TIMEOUT_MS = 45_000
LOAD_ATTEMPTS = 4                       # tries to (re)load the form before giving up
MAX_FORM_AGE_S = 10 * 60                # reload a form that has sat idle longer than this

# ---------------------------------------------------------------- reply classification
# OK    = server accepted it            ("Attendance marked successfully", "Already marked")
# FAIL  = server answered but rejected  ("Incorrect OTP.", "Invalid UG number")
# ERROR = never got an answer           (network down, timeout) - safe to retry
# CHECK = server answered something we don't recognise - look at it yourself
FAIL_RE = re.compile(r"\b(incorrect|invalid|wrong|expired|error|not|fail\w*|denied|closed|"
                     r"unauthori[sz]ed|exceed\w*|limit)\b|please enter", re.I)
OK_RE = re.compile(r"\b(success\w*|marked|submitted|recorded|saved|already|present|done)\b", re.I)


def classify(reached_server: bool, msg: str) -> str:
    if not reached_server:
        return "ERROR"
    if FAIL_RE.search(msg):
        return "FAIL"
    if OK_RE.search(msg):
        return "OK"
    return "CHECK"


# ---------------------------------------------------------------- input
UG_RE = re.compile(r"^UG/\d+/\d+$")   # UG/26/12345


def _norm(v) -> str:
    if v is None:
        return ""
    s = re.sub(r"\s+", "", str(v)).upper()   # " ug / 26 / 12345 " -> "UG/26/12345"
    if s.endswith(".0") and s[:-2].isdigit():   # Excel turns numbers into floats
        s = s[:-2]
    return s


def _read_rows(path: Path) -> list[list]:
    ext = path.suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        ws = load_workbook(path, read_only=True, data_only=True).active
        return [list(r) for r in ws.iter_rows(values_only=True) if r]
    if ext == ".csv":
        with path.open(newline="", encoding="utf-8-sig") as f:
            return [r for r in csv.reader(f) if r]
    return [[line] for line in path.read_text(encoding="utf-8-sig").splitlines()]


def _best_column(rows, test) -> int | None:
    counts: dict[int, int] = {}
    for r in rows:
        for i, v in enumerate(r):
            if test(_norm(v)):
                counts[i] = counts.get(i, 0) + 1
    return max(counts, key=counts.get) if counts else None


def _is_header(v) -> bool:
    s = _norm(v)
    return "UG" in s and not any(c.isdigit() for c in s)


def load_students(path: Path) -> tuple[list[tuple[str, str]], bool]:
    """Return ([(ug, name)], bare). Finds the UG column automatically:
    first a column of full IDs like UG/26/12345, otherwise the column headed
    "UG ..." or the column of plain numbers like 20292 (bare=True).
    The name is the first other text cell on the same row."""
    rows = _read_rows(path)

    ug_col = _best_column(rows, UG_RE.match)
    bare = ug_col is None
    if bare:
        header_col = next((i for r in rows[:3] for i, v in enumerate(r) if _is_header(v)), None)
        ug_col = header_col if header_col is not None else _best_column(rows, str.isdigit)
    if ug_col is None:
        sys.exit("Couldn't find a UG column (expected values like UG/26/12345 or 20292).")
    valid = str.isdigit if bare else UG_RE.match

    seen, out = set(), []
    for r in rows:
        raw = r[ug_col] if ug_col < len(r) else None
        ug = _norm(raw)
        if not ug or _is_header(raw):          # blank cell or header row
            continue
        if not valid(ug):
            print(f"  ! skipping row, UG value looks wrong: {raw!r}")
            continue
        if ug in seen:
            print(f"  ! skipping duplicate: {ug}")
            continue
        seen.add(ug)
        name = next((str(v).strip() for i, v in enumerate(r)
                     if i != ug_col and v not in (None, "") and not _norm(v).isdigit()), "")
        out.append((ug, name))
    return out, bare


# ---------------------------------------------------------------- browser session
def find_form_frame(page):
    """The form lives two iframes deep (…googleusercontent.com/blank)."""
    deadline = time.time() + LOAD_TIMEOUT_MS / 1000
    while time.time() < deadline:
        for f in page.frames:
            try:
                if f.query_selector("#ug") and f.query_selector("#otp"):
                    return f
            except Exception:
                pass
        page.wait_for_timeout(100)
    raise RuntimeError("form did not appear (couldn't find #ug / #otp)")


HEALTH_JS = f"() => !!(window.google && google.script && google.script.run && google.script.run.{SERVER_FN})"


class Session:
    """Owns the browser + a loaded form, and repairs itself when something breaks.
    Must only be used from the thread that created it (Playwright rule)."""

    def __init__(self, pw, headed: bool, log):
        self.pw, self.headed, self.log = pw, headed, log
        self.browser = self.context = self.page = self.frame = None
        self.loaded_at = 0.0

    def _launch(self):
        self.close()
        self.browser = self.pw.chromium.launch(headless=not self.headed)
        self.context = self.browser.new_context()

    def _load_once(self):
        if self.browser is None or not self.browser.is_connected():
            self._launch()
        if self.page is not None:
            try:
                self.page.close()
            except Exception:
                pass
        self.page = self.context.new_page()
        self.page.goto(URL, wait_until="domcontentloaded", timeout=LOAD_TIMEOUT_MS)
        self.frame = find_form_frame(self.page)
        self.frame.wait_for_function(HEALTH_JS, timeout=LOAD_TIMEOUT_MS)
        self.loaded_at = time.time()

    def healthy(self) -> bool:
        try:
            return bool(self.frame and self.frame.evaluate(HEALTH_JS))
        except Exception:
            return False

    def ensure_ready(self):
        """Make sure a working form is loaded; reload / relaunch as needed."""
        age = time.time() - self.loaded_at
        if self.healthy() and age < MAX_FORM_AGE_S:
            return
        if self.frame is not None and age >= MAX_FORM_AGE_S:
            self.log(f"(form has been open {age / 60:.0f} min - reloading a fresh copy first)")
        last = None
        for attempt in range(1, LOAD_ATTEMPTS + 1):
            try:
                if attempt == 3:          # two plain reloads failed: start a fresh browser
                    self._launch()
                self._load_once()
                if attempt > 1:
                    self.log(f"(form loaded on attempt {attempt})")
                return
            except Exception as e:
                last = e
                if attempt < LOAD_ATTEMPTS:
                    self.log(f"(form load attempt {attempt} failed: {short(e)} - retrying)")
                    time.sleep(2 * attempt)
        raise RuntimeError(f"could not load the form after {LOAD_ATTEMPTS} attempts: {short(last)}")

    def close(self):
        for obj in (self.browser,):
            try:
                if obj:
                    obj.close()
            except Exception:
                pass
        self.browser = self.context = self.page = self.frame = None


def short(e) -> str:
    s = str(e).strip().splitlines()[0] if str(e).strip() else e.__class__.__name__
    return s[:120]


# ---------------------------------------------------------------- mode: rpc (default)
# Calls the same server function the button calls, several at a time, from one page.
RPC_JS = """
async ({ugs, otp, fn, concurrency, retries, timeoutMs}) => {
  const call = (ug) => new Promise((resolve) => {
    const t = setTimeout(() => resolve({reached:false, msg:"no reply (timeout)"}), timeoutMs);
    try {
      google.script.run
        .withSuccessHandler(r => { clearTimeout(t); resolve({reached:true,  msg:String(r)}); })
        .withFailureHandler(e => { clearTimeout(t); resolve({reached:false, msg:String(e && e.message || e)}); })
        [fn](ug, otp);
    } catch (e) { clearTimeout(t); resolve({reached:false, msg:String(e && e.message || e)}); }
  });
  const results = new Array(ugs.length);
  let next = 0;
  async function worker() {
    while (next < ugs.length) {
      const i = next++;
      let r, attempt = 0;
      do {
        r = await call(ugs[i]);
        attempt++;
        if (!r.reached && attempt <= retries) await new Promise(s => setTimeout(s, 1000 * attempt));
      } while (!r.reached && attempt <= retries);
      results[i] = {ug: ugs[i], reached: r.reached, msg: r.msg, attempts: attempt};
    }
  }
  await Promise.all(Array.from({length: Math.min(concurrency, ugs.length)}, worker));
  return results;
}
"""


def run_rpc(session: Session, ugs, otp, concurrency, retries):
    session.ensure_ready()
    return session.frame.evaluate(RPC_JS, {
        "ugs": ugs, "otp": otp, "fn": SERVER_FN,
        "concurrency": concurrency, "retries": retries, "timeoutMs": CALL_TIMEOUT_MS,
    })


# ---------------------------------------------------------------- mode: click (fallback)
# Literally fills the two boxes and clicks "Mark Attendance", one student at a time.
def submit_by_click(frame, ug, otp):
    frame.fill("#ug", ug)
    frame.fill("#otp", otp)
    frame.evaluate("() => document.getElementById('message').innerText = ''")
    frame.click("button")
    frame.wait_for_function(
        "() => { const t = document.getElementById('message').innerText.trim();"
        "        return t && t !== 'Submitting...'; }",
        timeout=CALL_TIMEOUT_MS)
    return frame.inner_text("#message").strip()


def run_click(session: Session, ugs, otp, retries):
    results = []
    for ug in ugs:
        attempt, reached, msg = 0, False, ""
        while attempt <= retries and not reached:
            attempt += 1
            try:
                session.ensure_ready()
                msg, reached = submit_by_click(session.frame, ug, otp), True
            except Exception as e:
                msg = f"no reply ({short(e)})"
                session.frame = None        # force a reload before the next try
        results.append({"ug": ug, "reached": reached, "msg": msg, "attempts": attempt})
    return results


# ---------------------------------------------------------------- browser thread
def browser_worker(args, jobs: queue.Queue, replies: queue.Queue, ready: threading.Event,
                   load_error: list):
    """Runs on its own thread: loads the form immediately, then serves submit jobs."""
    log = lambda m: print(f"\n{m}", flush=True)
    try:
        with sync_playwright() as pw:
            session = Session(pw, args.headed, log)
            try:
                session.ensure_ready()
            except Exception as e:
                load_error.append(e)          # main thread will try again on first job
            finally:
                ready.set()

            while (job := jobs.get()) is not None:
                otp, batch = job
                try:
                    if args.mode == "rpc":
                        res = run_rpc(session, batch, otp, args.concurrency, args.retries)
                        # Anything that never reached the server: reload the form and try once more.
                        lost = [r["ug"] for r in res if not r["reached"]]
                        if lost:
                            log(f"({len(lost)} got no reply - reloading form and retrying them)")
                            session.frame = None
                            again = {r["ug"]: r for r in
                                     run_rpc(session, lost, otp, args.concurrency, args.retries)}
                            res = [again.get(r["ug"], r) for r in res]
                    else:
                        res = run_click(session, batch, otp, args.retries)
                    replies.put(res)
                except Exception as e:
                    # Never crash the run: report everyone in this batch as ERROR.
                    replies.put([{"ug": u, "reached": False, "msg": f"no reply ({short(e)})",
                                  "attempts": 0} for u in batch])
            session.close()
    except Exception as e:                    # e.g. Playwright/Chromium missing
        load_error.append(e)
        ready.set()
        while (job := jobs.get()) is not None:
            replies.put([{"ug": u, "reached": False, "msg": f"browser failed: {short(e)}",
                          "attempts": 0} for u in job[1]])


# ---------------------------------------------------------------- main
MARK = {"OK": "OK   ", "FAIL": "FAIL ", "ERROR": "ERROR", "CHECK": "CHECK"}


def save_results(final: dict, ugs: list) -> Path | None:
    rows = [final[u] for u in ugs if u in final]
    if not rows:
        return None
    for attempt in range(2):
        out = Path(f"results_{time.strftime('%Y%m%d_%H%M%S')}{'_b' if attempt else ''}.csv")
        try:
            with out.open("w", newline="", encoding="utf-8") as f:
                w = csv.DictWriter(f, fieldnames=["ug", "name", "status", "msg", "otp", "attempts"],
                                   extrasaction="ignore")
                w.writeheader()
                w.writerows(rows)
            return out
        except OSError:
            continue
    print("! could not write the results file")
    return None


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("file", type=Path, help=".xlsx / .csv / .txt with names + UG numbers (UG column auto-detected)")
    ap.add_argument("--otp", help="skip the prompt and use this OTP (single round, no retry prompt)")
    ap.add_argument("--prefix", default="UG/26/", help='added to plain numbers (default "UG/26/")')
    ap.add_argument("--mode", choices=["rpc", "click"], default="rpc")
    ap.add_argument("--concurrency", type=int, default=4, help="parallel submissions (rpc mode)")
    ap.add_argument("--retries", type=int, default=2, help="extra tries per student when there's no reply")
    ap.add_argument("--headed", action="store_true", help="show the browser (debugging)")
    args = ap.parse_args()

    # ---- read the sheet
    if not args.file.exists():
        sys.exit(f"File not found: {args.file}")
    try:
        students, bare = load_students(args.file)
    except PermissionError:
        sys.exit(f"Can't open {args.file.name} - close it in Excel and try again.")
    except Exception as e:
        sys.exit(f"Can't read {args.file.name}: {short(e)}")
    if students and bare:
        prefix = re.sub(r"\s+", "", args.prefix).upper()
        if prefix and not prefix.endswith("/"):
            prefix += "/"
        students = [(prefix + u, n) for u, n in students]
    ugs = [u for u, _ in students]
    names = dict(students)
    if not ugs:
        sys.exit("No UG numbers found in the sheet.")
    print(f"Loaded {len(ugs)} student(s):")
    for i, (u, n) in enumerate(students, 1):
        print(f"  {i:>3}. {u:<16} {n}")

    # ---- start loading the form in the background
    jobs: queue.Queue = queue.Queue()
    replies: queue.Queue = queue.Queue()
    ready = threading.Event()
    load_error: list = []
    worker = threading.Thread(target=browser_worker, args=(args, jobs, replies, ready, load_error),
                              daemon=True)
    worker.start()
    print("\n(Form is loading in the background - type the OTP whenever you're ready.)")

    final: dict[str, dict] = {}
    pending = list(ugs)
    otp = (args.otp or "").strip()
    last_otp = ""
    try:
        while pending:
            if not otp:
                otp = input("\nEnter OTP and press Enter: ").strip()
                if not otp:
                    print("No OTP given - stopping.")
                    break
            if not ready.is_set():
                print("Finishing form load...", flush=True)
                ready.wait()
            if load_error and len(load_error) == 1:
                print(f"(first form load failed: {short(load_error[0])} - will retry while submitting)")
                load_error.append(None)      # only say it once

            print(f"Submitting {len(pending)} student(s)...", flush=True)
            t0 = time.time()
            jobs.put((otp, pending))
            results = replies.get()
            took = time.time() - t0

            for r in results:
                r["status"] = classify(r["reached"], r["msg"])
                r["name"] = names.get(r["ug"], "")
                r["otp"] = otp
                final[r["ug"]] = r
                print(f"  {MARK[r['status']]}  {r['ug']:<16} {r['name']:<20} {r['msg']}")
            last_otp = otp

            counts = {s: sum(1 for u in ugs if u in final and final[u]["status"] == s) for s in MARK}
            print(f"\n  OK {counts['OK']}  |  FAIL {counts['FAIL']}  |  ERROR {counts['ERROR']}"
                  f"  |  CHECK {counts['CHECK']}   ({took:.1f}s)")

            # Retry FAIL (e.g. wrong OTP) and ERROR (no reply). CHECK is left for you to look at.
            pending = [r["ug"] for r in results if r["status"] in ("FAIL", "ERROR")]
            if not pending or args.otp:
                break
            ans = input(f"\nRetry the {len(pending)} FAIL/ERROR? "
                        f"Type a new OTP, Enter = same OTP, q = quit: ").strip()
            if ans.lower() == "q":
                break
            otp = ans or last_otp
    except (KeyboardInterrupt, EOFError):
        print("\nStopped.")
    finally:
        jobs.put(None)

    out = save_results(final, ugs)
    if final:
        bad = [final[u] for u in ugs if u in final and final[u]["status"] != "OK"]
        not_sent = [u for u in ugs if u not in final]
        print(f"\nFinal: {len(final) - len(bad)} OK, {len(bad)} not OK"
              + (f", {len(not_sent)} never sent" if not_sent else "") + ".")
        for r in bad:
            print(f"  {MARK[r['status']]}  {r['ug']} ({r['name']}): {r['msg']}")
        if out:
            print(f"Saved to {out}")
        sys.exit(0 if not bad and not not_sent else 1)


if __name__ == "__main__":
    main()
